"""Content acquisition and locatable normalization. No model or ASR is embedded."""
import argparse
import hashlib
import importlib.util
import json
import re
import sqlite3
import sys
import urllib.parse
import urllib.request
import zipfile
import xml.etree.ElementTree as ET
from html.parser import HTMLParser
from pathlib import Path

from bundle_tool import digest, read, require, save

VERSION = "ayou-source-1.1.1"
MAX_BYTES = 32 * 1024 * 1024
MEDIA = {".mp4", ".mkv", ".mov", ".webm", ".mp3", ".wav", ".m4a", ".ogg"}


class TextHTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.skip = 0
        self.parts = []
        self.title_parts = []
        self.in_title = False

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "nav", "footer", "noscript"}:
            self.skip += 1
        if tag == "title":
            self.in_title = True
        if not self.skip and tag in {"p", "br", "li", "div", "article", "section", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")
            if tag in {"h1", "h2", "h3", "h4"}:
                self.parts.append("#" * int(tag[1]) + " ")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "nav", "footer", "noscript"}:
            self.skip = max(0, self.skip - 1)
        if tag == "title":
            self.in_title = False
        if tag in {"p", "li", "div", "article", "section", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_data(self, data):
        if self.in_title:
            self.title_parts.append(data)
        elif not self.skip:
            self.parts.append(data)

    def text(self):
        return re.sub(r"\n\s*\n+", "\n\n", "".join(self.parts)).strip()


def fetch(url):
    require(urllib.parse.urlparse(url).scheme in {"http", "https"}, "Only HTTP(S) source URLs")
    request = urllib.request.Request(url, headers={"User-Agent": "ayou-distill/1.1"})
    with urllib.request.urlopen(request, timeout=30) as response:
        if response.headers.get_content_type().startswith(("video/", "audio/")):
            return b"", response.headers.get_content_type(), response.geturl()
        raw = response.read(MAX_BYTES + 1)
        require(len(raw) <= MAX_BYTES, "Source exceeds 32 MiB; provide a local file")
        return raw, response.headers.get_content_type(), response.geturl()


def text_decode(raw):
    for encoding in ("utf-8-sig", "utf-16", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeError:
            continue
    raise ValueError("Cannot decode text; supply encoding-normalized UTF-8")


def html_text(raw):
    parser = TextHTML()
    parser.feed(text_decode(raw))
    return parser.text(), "".join(parser.title_parts).strip()


def milliseconds(text):
    parts = text.strip().replace(",", ".").split(":")
    require(len(parts) in {2, 3}, "Invalid subtitle timestamp")
    value = float(parts[-1]) + 60 * int(parts[-2])
    if len(parts) == 3:
        value += 3600 * int(parts[0])
    return round(value * 1000)


def subtitle_records(text):
    records = []
    lines = text.replace("\r\n", "\n").splitlines()
    index = 0
    while index < len(lines):
        match = re.match(r"\s*([\d:.,]+)\s*-->\s*([\d:.,]+)", lines[index])
        if not match:
            index += 1
            continue
        start, end = milliseconds(match[1]), milliseconds(match[2])
        require(end >= start >= 0, "Reversed/negative subtitle timestamps")
        index += 1
        body = []
        while index < len(lines) and lines[index].strip():
            body.append(re.sub(r"<[^>]+>", "", lines[index]))
            index += 1
        text_body = " ".join(body).strip()
        if text_body:
            records.append(dict(text=text_body, location=f"{start/1000:.3f}-{end/1000:.3f}s", time_start_ms=start, time_end_ms=end, heading_path=[]))
    require(records, "No subtitle cues found")
    return records


def text_records(text, prefix=""):
    records, headings = [], []
    for line_no, line in enumerate(text.splitlines(), 1):
        if not line.strip():
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)", line)
        if heading:
            level, title = len(heading[1]), heading[2]
            headings = headings[:level-1] + [title]
        records.append(dict(text=line, location=f"{prefix}line:{line_no}", heading_path=list(headings)))
    require(records, "No readable text")
    return records


def parse_local(path):
    path = Path(path).resolve()
    ext, raw = path.suffix.lower(), path.read_bytes()
    if ext in MEDIA:
        return [], {"status": "needs_transcript", "reason": "Local audio/video needs a transcript provider; no built-in ASR"}, raw
    if ext == ".pdf":
        from pypdf import PdfReader
        records, empty = [], []
        for page_no, page in enumerate(PdfReader(path).pages, 1):
            text = page.extract_text() or ""
            if text.strip():
                records += text_records(text, f"reader-page:{page_no}/")
            else:
                empty.append(page_no)
        return records, {"status": "needs_ocr" if not records else "needs_review", "empty_pages": empty, "reason": "PDF text extraction does not verify figures, formulas or OCR"}, raw
    if ext == ".docx":
        from docx import Document
        document = Document(path)
        records = [dict(text=p.text, location=f"paragraph:{i}", heading_path=[]) for i,p in enumerate(document.paragraphs, 1) if p.text.strip()]
        for ti, table in enumerate(document.tables, 1):
            for ri, row in enumerate(table.rows, 1):
                records.append(dict(text=" | ".join(cell.text for cell in row.cells), location=f"table:{ti}/row:{ri}", heading_path=[]))
        return records, {"status": "needs_review", "reason": "Paragraph and table positions retained; inline figures not inspected"}, raw
    if ext == ".epub":
        records = []
        with zipfile.ZipFile(path) as archive:
            container = ET.fromstring(archive.read("META-INF/container.xml"))
            package = next(x.attrib["full-path"] for x in container.iter() if x.tag.endswith("rootfile"))
            opf = ET.fromstring(archive.read(package))
            items = {x.attrib["id"]: x.attrib["href"] for x in opf.iter() if x.tag.endswith("item")}
            import posixpath
            for item in (x for x in opf.iter() if x.tag.endswith("itemref")):
                name = posixpath.normpath(posixpath.join(posixpath.dirname(package), urllib.parse.unquote(items[item.attrib["idref"]])))
                require(not name.startswith("../"), "EPUB spine escapes archive")
                info = archive.getinfo(name)
                require(info.file_size <= MAX_BYTES, "EPUB chapter too large")
                text, title = html_text(archive.read(name))
                if text.strip():
                    records += text_records(text, f"epub:{name}/")
        return records, {"status": "needs_review", "reason": "EPUB spine order retained; figures/footnote relations need review"}, raw
    text = text_decode(raw)
    if ext in {".srt", ".vtt"}:
        return subtitle_records(text), {"status": "needs_review", "reason": "Subtitle accuracy and visual dependencies not verified"}, raw
    if ext == ".json":
        data = json.loads(text)
        records = []
        for cue in data.get("body", data.get("segments", [])):
            start, end = float(cue.get("from", cue.get("start", 0))), float(cue.get("to", cue.get("end", 0)))
            require(end >= start >= 0, "Invalid JSON transcript timing")
            body = cue.get("content", cue.get("text", "")).strip()
            if body:
                records.append(dict(text=body, location=f"{start:.3f}-{end:.3f}s", time_start_ms=round(start*1000), time_end_ms=round(end*1000), heading_path=[]))
        require(records, "JSON must contain timed body or segments; other JSON requires explicit conversion")
        return records, {"status": "needs_review"}, raw
    if ext in {".html", ".htm"}:
        text, title = html_text(raw)
        return text_records(text), {"status": "needs_review", "title": title, "reason": "Extracted HTML needs main-content verification"}, raw
    require(ext in {".txt", ".md", ".markdown"}, "Unsupported local format")
    return text_records(text), {"status": "needs_review"}, raw


def media_url(url, languages):
    try:
        import yt_dlp
    except ImportError:
        return [], {"status": "needs_dependency", "reason": "Install yt-dlp for video metadata/subtitles"}, url.encode()
    with yt_dlp.YoutubeDL({"skip_download": True, "quiet": True, "no_warnings": True, "noplaylist": True, "socket_timeout": 25, "retries": 1}) as downloader:
        info = downloader.extract_info(url, download=False)
    require(info and not info.get("entries"), "Playlist: select an episode or process entries separately")
    metadata = {"title": info.get("title"), "author": info.get("uploader"), "published": info.get("upload_date"), "duration": info.get("duration"), "status": "needs_transcript", "reason": "No preferred text subtitle available; use a transcript provider"}
    for key, automatic in (("subtitles", False), ("automatic_captions", True)):
        captions = info.get(key) or {}
        for preferred in languages:
            for language in sorted(captions):
                if language != preferred and not language.startswith(preferred + "-"):
                    continue
                for choice in sorted(captions[language], key=lambda x: x.get("ext") != "vtt"):
                    if choice.get("ext") not in {"vtt", "srt"}:
                        continue
                    raw, _, _ = fetch(choice["url"])
                    records = subtitle_records(text_decode(raw))
                    metadata.update(status="needs_review", subtitle_language=language, automatic=automatic, reason="Caption text acquired; accuracy/visual content not verified")
                    return records, metadata, raw
    return [], metadata, json.dumps(metadata, ensure_ascii=False).encode()


def acquire(source, root, languages):
    if not re.match(r"https?://", source):
        records, metadata, raw = parse_local(source)
        return records, metadata, raw
    host = urllib.parse.urlparse(source).hostname or ""
    if any(host == h or host.endswith("." + h) for h in ("youtube.com", "youtu.be", "bilibili.com", "b23.tv", "vimeo.com", "douyin.com")):
        return media_url(source, languages)
    raw, mime, final = fetch(source)
    ext = Path(urllib.parse.urlparse(final).path).suffix.lower()
    if mime.startswith(("video/", "audio/")):
        return [], {"status":"needs_transcript", "reason":"Direct media URL requires a transcript provider", "final_url":final}, raw
    if mime in {"application/rss+xml", "application/atom+xml", "text/xml", "application/xml"}:
        tree = ET.fromstring(raw)
        links = [x.attrib.get("url") for x in tree.iter() if x.tag.endswith("transcript") and x.attrib.get("url")]
        return [], {"status":"needs_episode_selection", "transcript_urls":links, "reason":"Select a podcast episode and its transcript; feed summaries are not full content"}, raw
    if ext in {".md", ".txt", ".srt", ".vtt", ".pdf", ".epub", ".docx", ".json"}:
        filename = hashlib.sha256(raw).hexdigest() + ext
        target = root / "raw" / filename
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(raw)
        records, metadata, _ = parse_local(target)
    else:
        text, title = html_text(raw)
        records = text_records(text) if text.strip() else []
        metadata = {"status":"needs_review" if len(text) >= 200 else "needs_browser", "title":title, "reason":"Verify extracted main text; login/JavaScript pages may need browser access"}
    metadata["final_url"] = final
    return records, metadata, raw


def ingest(source, out, max_chars=4000, languages=("zh", "en")):
    require(max_chars >= 100, "max_chars must be >=100")
    root = Path(out).resolve()
    root.mkdir(parents=True, exist_ok=True)
    records, metadata, raw = acquire(source, root, languages)
    source_hash = hashlib.sha256(raw).hexdigest()
    source_id = "src-" + hashlib.sha256(str(source).encode()).hexdigest()[:12]
    version = hashlib.sha256((VERSION + source_hash + str(max_chars)).encode()).hexdigest()
    version_root = root / "sources" / source_id / version
    if version_root.exists():
        manifest=read(version_root/"manifest.json")
        require(all(digest(version_root/name)==sha for name,sha in manifest["files"].items()),"Cached preprocessing edited; preserve and use a new output directory")
        cached=read(version_root/"document.json")
        return {"directory":str(version_root),"chunks":len(read(version_root/"chunks.json")),"status":cached["metadata"]["status"],"metadata":cached["metadata"],"cached":True}
    document = {"parser":VERSION, "source_id":source_id, "version":version, "origin":str(source), "source_sha256":source_hash, "metadata":metadata, "records":records}
    chunks, pending = [], []
    def flush():
        nonlocal pending
        if not pending:
            return
        body="\n".join(x["text"] for x in pending)
        chunk={**pending[0],"text":body,"locations":[x["location"] for x in pending],"seq":len(chunks),"source_id":source_id,"version":version,"chunk_id":f"{source_id}-{len(chunks):06d}","content_hash":hashlib.sha256(body.encode()).hexdigest()}
        if all("time_start_ms" in x for x in pending):
            chunk["time_start_ms"]=min(x["time_start_ms"] for x in pending)
            chunk["time_end_ms"]=max(x["time_end_ms"] for x in pending)
        chunks.append(chunk)
        pending=[]
    for record in records:
        # Splits retain page/time/line evidence; no timing interpolation is invented.
        for offset in range(0, len(record["text"]), max_chars):
            body = record["text"][offset:offset+max_chars]
            element={**record,"text":body,"char_offset":offset}
            if pending:
                old_page=pending[0]["location"].split("/line:")[0] if pending[0]["location"].startswith("reader-page:") else None
                new_page=record["location"].split("/line:")[0] if record["location"].startswith("reader-page:") else None
                if pending[0]["heading_path"]!=record["heading_path"] or old_page!=new_page or sum(len(x["text"])+1 for x in pending)+len(body)>max_chars:
                    flush()
            pending.append(element)
    flush()
    save(version_root / "document.json", document)
    save(version_root / "chunks.json", chunks)
    raw_path = version_root / "source.bin"
    raw_path.write_bytes(raw)
    save(version_root / "manifest.json", {"source_id":source_id,"version":version,"source_sha256":source_hash,"files":{name:digest(version_root/name) for name in ("document.json","chunks.json","source.bin")}})
    return {"directory":str(version_root),"chunks":len(chunks),"status":metadata["status"],"metadata":metadata}


def query(chunks_path, words, limit=5, neighbors=1):
    require(words.strip() and 1 <= limit <= 100 and 0 <= neighbors <= 10, "Invalid query options")
    chunks = read(chunks_path)
    def tokenize(text):
        return re.findall(r"[a-zA-Z0-9]+", text.lower()) + [s[i:i+2] for s in re.findall(r"[\u4e00-\u9fff]+",text) for i in range(max(1,len(s)-1))]
    connection=sqlite3.connect(":memory:")
    connection.execute("CREATE VIRTUAL TABLE search USING fts5(body)")
    for chunk in chunks:
        connection.execute("INSERT INTO search(body) VALUES(?)",(" ".join(tokenize(chunk["text"])),))
    tokens=tokenize(words)
    require(tokens,"Query has no searchable terms")
    expression=" OR ".join('"'+t+'"' for t in tokens)
    rows=connection.execute("SELECT rowid FROM search WHERE search MATCH ? ORDER BY rank LIMIT ?",(expression,limit)).fetchall()
    chosen={n for row in rows for n in range(max(0,row[0]-1-neighbors),min(len(chunks),row[0]+neighbors))}
    connection.close()
    return [chunks[n] for n in sorted(chosen)]


def doctor():
    return {"python":sys.version.split()[0],"sqlite_fts5":sqlite3.connect(":memory:").execute("SELECT sqlite_compileoption_used('ENABLE_FTS5')").fetchone()[0] == 1,"dependencies":{name:bool(importlib.util.find_spec(name)) for name in ("yt_dlp","pypdf","docx")},"built_in_asr":False,"ocr":False}


def main():
    parser=argparse.ArgumentParser()
    sub=parser.add_subparsers(dest="command",required=True)
    sub.add_parser("doctor")
    p=sub.add_parser("ingest"); p.add_argument("source"); p.add_argument("--out",required=True); p.add_argument("--max-chars",type=int,default=4000); p.add_argument("--languages",default="zh,en")
    p=sub.add_parser("collection"); p.add_argument("sources",nargs="+"); p.add_argument("--out",required=True)
    p=sub.add_parser("query"); p.add_argument("chunks"); p.add_argument("words"); p.add_argument("--limit",type=int,default=5); p.add_argument("--neighbors",type=int,default=1)
    args=parser.parse_args()
    try:
        if args.command == "doctor": result=doctor()
        elif args.command == "query": result=query(args.chunks,args.words,args.limit,args.neighbors)
        elif args.command == "collection":
            result=[]
            for source in args.sources:
                try: result.append({"source":source,**ingest(source,args.out)})
                except Exception as error: result.append({"source":source,"status":"acquisition_failed","reason":str(error)})
            save(Path(args.out)/"collection.json",result)
        else: result=ingest(args.source,args.out,args.max_chars,args.languages.split(","))
        print(json.dumps(result,ensure_ascii=False,indent=2))
        return 0
    except Exception as error:
        print(json.dumps({"status":"acquisition_failed","reason":str(error)},ensure_ascii=False))
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
